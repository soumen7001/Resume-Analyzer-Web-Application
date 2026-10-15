import requests
import json
import io
from docx import Document

BASE_URL = "http://localhost:5000"

def test_health():
    print("\n--- 1. Testing /health ---")
    r = requests.get(f"{BASE_URL}/health")
    assert r.status_code == 200, f"Health check failed: {r.status_code}"
    data = r.json()
    print("Response:", json.dumps(data, indent=2))
    assert data["status"] == "healthy"
    print("PASSED: /health is healthy.")

def test_skills_database():
    print("\n--- 2. Testing /api/skills ---")
    r = requests.get(f"{BASE_URL}/api/skills")
    assert r.status_code == 200
    data = r.json()
    tech_count = len(data["data"]["technical"])
    soft_count = len(data["data"]["soft"])
    print(f"Technical skills: {tech_count}, Soft skills: {soft_count}")
    assert tech_count > 50 and soft_count > 10
    print("PASSED: /api/skills returned database.")

def test_analyze_sample():
    print("\n--- 3. Testing /api/analyze-sample ---")
    r = requests.get(f"{BASE_URL}/api/analyze-sample")
    assert r.status_code == 200
    data = r.json()
    assert data["success"] is True
    res = data["data"]
    print(f"Sample Overall Score: {res.get('overall_score')}, ATS Score: {res.get('ats_score')}")
    print(f"Detected Skills: {len(res.get('detected_skills', {}).get('technical', []))} tech skills")
    assert res.get('overall_score', 0) > 50
    print("PASSED: /api/analyze-sample working perfectly.")

def test_upload_and_analyze_docx():
    print("\n--- 4. Testing /upload-resume with generated DOCX ---")
    doc = Document()
    doc.add_heading("Jane Doe", 0)
    doc.add_paragraph("Email: jane.doe@example.com | Phone: 123-456-7890 | LinkedIn: linkedin.com/in/janedoe")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Senior Full Stack Software Engineer with 6+ years of experience in Python, Django, React, AWS, Docker, and PostgreSQL.")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Architect at ABC Tech (2020 - Present)")
    doc.add_paragraph("Spearheaded microservice migration reducing latency by 45% using FastAPI and Kubernetes.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, Django, FastAPI, React, TypeScript, Docker, Kubernetes, AWS, PostgreSQL, Git, CI/CD, Agile")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Computer Engineering, University of Technology, 2018")

    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    docx_bytes.seek(0)

    files = {"resume": ("jane_doe_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r = requests.post(f"{BASE_URL}/upload-resume", files=files)
    assert r.status_code == 200, f"Upload failed: {r.text}"
    data = r.json()
    assert data["success"] is True
    parsed = data["data"]
    print(f"Parsed Contact: {parsed.get('contact_info')}")
    print(f"Parsed Skills: {parsed.get('skills')}")
    assert parsed.get("contact_info", {}).get("email") == "jane.doe@example.com"
    print("PASSED: DOCX parsing and text extraction.")
    return parsed

def test_analyze_with_job_matching(parsed_resume):
    print("\n--- 5. Testing /analyze-resume with Job Description ---")
    jd = """
    We are seeking a Senior Python Engineer experienced with Python, FastAPI, Docker, Kubernetes,
    AWS, and PostgreSQL. Must have strong leadership and agile experience. React is a bonus.
    """
    payload = {
        "resume_text": parsed_resume["raw_text"],
        "job_description": jd
    }
    r = requests.post(f"{BASE_URL}/analyze-resume", json=payload)
    assert r.status_code == 200, f"Analysis failed: {r.text}"
    res = r.json()["data"]
    print(f"Overall Score: {res.get('overall_score')}")
    print(f"Job Match Score: {res.get('job_match_score')}")
    print(f"Matching Keywords: {res.get('job_match', {}).get('matching_keywords')}")
    print(f"Missing Keywords: {res.get('job_match', {}).get('missing_keywords')}")
    assert res.get('overall_score') is not None
    print("PASSED: Full ATS scoring and TF-IDF Job Matching.")

def test_resume_generation_and_export(parsed_resume):
    print("\n--- 6. Testing /generate-resume & /export-resume ---")
    gen_payload = {
        "parsed_resume": parsed_resume,
        "job_description": "Senior Python Engineer",
        "target_role": "Lead Software Engineer"
    }
    r = requests.post(f"{BASE_URL}/generate-resume", json=gen_payload)
    assert r.status_code == 200, f"Generation failed: {r.text}"
    gen_data = r.json()["data"]
    print(f"Generated Summary: {gen_data.get('summary')}")
    assert gen_data.get("summary") != ""

    export_payload = {
        "resume_data": gen_data,
        "template": "tech",
        "format": "docx"
    }
    r_exp = requests.post(f"{BASE_URL}/export-resume", json=export_payload)
    assert r_exp.status_code == 200
    assert len(r_exp.content) > 1000
    print(f"Exported DOCX size: {len(r_exp.content)} bytes")
    print("PASSED: Optimized Resume Generation and DOCX Export.")

if __name__ == "__main__":
    test_health()
    test_skills_database()
    test_analyze_sample()
    parsed = test_upload_and_analyze_docx()
    test_analyze_with_job_matching(parsed)
    test_resume_generation_and_export(parsed)
    print("\n==========================================")
    print("  ALL 6 BACKEND & PIPELINE TESTS PASSED!  ")
    print("==========================================")
